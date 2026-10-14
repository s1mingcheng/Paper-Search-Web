"""Paper Search Web - local paper retrieval and Excel export demo.

The application uses only public APIs by default. It does not bypass paywalls,
login walls, CAPTCHAs, or website anti-bot controls.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import ssl
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from dotenv import load_dotenv

from deepseek import DeepSeekError, DeepSeekService
from translator import (
    TranslationService,
    add_variants,
    all_variant_terms,
    build_flat_query,
    build_grouped_query,
    build_primary_query,
    group_variants,
)

APP_VERSION = "1.3.0"
BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env", override=False, encoding="utf-8-sig")
STATIC_DIR = BASE_DIR / "static"
DATA_DIR = BASE_DIR / "data"
JOURNAL_METRICS_FILE = DATA_DIR / "journal_metrics.csv"
TRANSLATION_DICTIONARY_FILE = DATA_DIR / "zh_en_terms.csv"
TRANSLATION_CACHE_FILE = DATA_DIR / "translation_cache.json"
DEEPSEEK_CACHE_FILE = DATA_DIR / "deepseek_summary_cache.json"
DEEPSEEK_TERMS_CACHE_FILE = DATA_DIR / "deepseek_terms_cache.json"
DEEPSEEK_CONFIG_FILE = DATA_DIR / "deepseek_config.json"
USER_AGENT = f"PaperSearchWeb/{APP_VERSION} (+http://127.0.0.1)"
DEFAULT_TIMEOUT = 30
MAX_RESULTS = 100
DEFAULT_MAX_RESULTS = 10


class SourceError(RuntimeError):
    """A recoverable error from one upstream source."""


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data:
            self.parts.append(data)

    def text(self) -> str:
        return " ".join(self.parts)


def strip_markup(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        value = " ".join(str(part) for part in value if part is not None)
    text = str(value)
    if "<" in text and ">" in text:
        parser = _TextExtractor()
        try:
            parser.feed(text)
            text = parser.text()
        except Exception:
            text = re.sub(r"<[^>]+>", " ", text)
    return unescape(text)


def clean_text(value: Any, limit: int | None = None) -> str:
    text = unicodedata.normalize("NFKC", strip_markup(value))
    text = re.sub(r"\s+", " ", text).strip()
    if limit and len(text) > limit:
        return text[: max(0, limit - 1)].rstrip() + "…"
    return text


def as_text(value: Any) -> str:
    return clean_text(value)


def first(value: Any, default: str = "") -> str:
    if isinstance(value, (list, tuple)):
        return as_text(value[0]) if value else default
    return as_text(value) if value is not None else default


def normalize_doi(value: Any) -> str:
    text = as_text(value).lower()
    text = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", text)
    text = re.sub(r"^doi:\s*", "", text)
    return text.strip()


def normalize_title(value: Any) -> str:
    text = unicodedata.normalize("NFKC", as_text(value)).casefold()
    return re.sub(r"[^\w\u3400-\u9fff]+", "", text)


def normalize_key(value: Any) -> str:
    text = unicodedata.normalize("NFKC", as_text(value)).casefold()
    return re.sub(r"\s+", " ", text).strip()


def parse_int(value: Any, default: int | None = None) -> int | None:
    try:
        if value is None or value == "":
            return default
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return default


def parse_float(value: Any) -> float | None:
    try:
        if value is None or value == "":
            return None
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def parse_year(value: Any) -> int | None:
    text = as_text(value)
    match = re.search(r"(19|20)\d{2}", text)
    return int(match.group(0)) if match else None


def safe_url(value: Any) -> str:
    url = as_text(value)
    if url.startswith("http://") or url.startswith("https://"):
        return url
    return ""


def doi_url(doi: str) -> str:
    return f"https://doi.org/{doi}" if doi else ""


def split_keywords(raw: Any) -> list[str]:
    if isinstance(raw, list):
        chunks: list[str] = []
        for item in raw:
            chunks.extend(split_keywords(item))
    else:
        text = as_text(raw)
        chunks = re.split(r"[,，;；\n\r\t]+", text)
    result: list[str] = []
    seen: set[str] = set()
    for chunk in chunks:
        keyword = clean_text(chunk, limit=120)
        key = normalize_key(keyword)
        if keyword and key not in seen:
            seen.add(key)
            result.append(keyword)
    return result


class JournalMetrics:
    """Load a local journal -> impact factor mapping.

    JCR impact factors are licensed data, so the project ships no JCR values.
    Users can place their licensed/exported values in data/journal_metrics.csv.
    """

    def __init__(self, path: Path) -> None:
        self.path = path
        self._mtime: float | None = None
        self._by_name: dict[str, tuple[float, str]] = {}
        self._by_issn: dict[str, tuple[float, str]] = {}

    def _load_if_needed(self) -> None:
        try:
            mtime = self.path.stat().st_mtime
        except OSError:
            mtime = -1.0
        if mtime == self._mtime:
            return
        by_name: dict[str, tuple[float, str]] = {}
        by_issn: dict[str, tuple[float, str]] = {}
        if mtime >= 0:
            try:
                with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
                    lines = [line for line in handle if not line.lstrip().startswith("#")]
                reader = csv.DictReader(io.StringIO("".join(lines)))
                for row in reader:
                    journal = clean_text(row.get("journal") or row.get("journal_name") or "")
                    impact = parse_float(row.get("impact_factor") or row.get("jif"))
                    if not journal or impact is None:
                        continue
                    year = clean_text(row.get("year") or "")
                    source = clean_text(row.get("source") or "")
                    source_label = "本地映射"
                    if year:
                        source_label += f" {year}"
                    if source:
                        source_label += f" · {source}"
                    if journal:
                        by_name[normalize_key(journal)] = (impact, source_label)
                    for key in ("issn", "eissn"):
                        for issn in re.split(r"[,，;；\s]+", clean_text(row.get(key) or "")):
                            issn_key = re.sub(r"[^0-9Xx]", "", issn).upper()
                            if issn_key:
                                by_issn[issn_key] = (impact, source_label)
            except Exception as exc:  # Keep search usable if the optional file is malformed.
                print(f"[warn] Could not read journal metrics file: {exc}", file=sys.stderr)
        self._mtime = mtime
        self._by_name = by_name
        self._by_issn = by_issn

    def lookup(self, journal: Any, issn: Any = "") -> tuple[float | None, str]:
        self._load_if_needed()
        journal_key = normalize_key(journal)
        if journal_key and journal_key in self._by_name:
            return self._by_name[journal_key]
        issn_text = clean_text(issn)
        for candidate in re.split(r"[,，;；\s]+", issn_text):
            issn_key = re.sub(r"[^0-9Xx]", "", candidate).upper()
            if issn_key and issn_key in self._by_issn:
                return self._by_issn[issn_key]
        return None, "未配置"

    def count(self) -> int:
        self._load_if_needed()
        return len(self._by_name)


JOURNAL_METRICS = JournalMetrics(JOURNAL_METRICS_FILE)
TRANSLATOR = TranslationService(TRANSLATION_DICTIONARY_FILE, TRANSLATION_CACHE_FILE)
DEEPSEEK = DeepSeekService(DEEPSEEK_CACHE_FILE, DEEPSEEK_CONFIG_FILE, DEEPSEEK_TERMS_CACHE_FILE)


def request_bytes(
    url: str,
    *,
    headers: dict[str, str] | None = None,
    timeout: int = DEFAULT_TIMEOUT,
) -> bytes:
    request_headers = {
        "User-Agent": USER_AGENT,
        "Accept": "*/*",
        "Accept-Encoding": "identity",
    }
    if headers:
        request_headers.update(headers)
    request = urllib.request.Request(url, headers=request_headers)
    try:
        context = ssl.create_default_context()
        with urllib.request.urlopen(request, timeout=timeout, context=context) as response:
            return response.read()
    except urllib.error.HTTPError as exc:
        detail = ""
        try:
            detail = exc.read(400).decode("utf-8", errors="replace")
        except Exception:
            pass
        detail = clean_text(detail, limit=180)
        raise SourceError(f"HTTP {exc.code} {exc.reason}{': ' + detail if detail else ''}") from exc
    except urllib.error.URLError as exc:
        raise SourceError(f"网络连接失败: {exc.reason}") from exc
    except TimeoutError as exc:
        raise SourceError("请求超时") from exc


def request_json(
    url: str,
    *,
    headers: dict[str, str] | None = None,
    timeout: int = DEFAULT_TIMEOUT,
) -> Any:
    raw = request_bytes(url, headers=headers, timeout=timeout)
    try:
        return json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise SourceError("上游返回了无法解析的 JSON") from exc


def build_boolean_query(keywords: list[str], match_mode: str, *, quote: bool = False) -> str:
    joiner = " AND " if match_mode == "all" else " OR "
    parts: list[str] = []
    for keyword in keywords:
        escaped = keyword.replace('"', '\\"')
        parts.append(f'"{escaped}"' if quote else escaped)
    return joiner.join(parts)


def clamp_max_results(value: Any) -> int:
    number = parse_int(value, DEFAULT_MAX_RESULTS) or DEFAULT_MAX_RESULTS
    return max(1, min(MAX_RESULTS, number))


def normalize_year_range(raw_from: Any, raw_to: Any) -> tuple[int | None, int | None]:
    start = parse_year(raw_from)
    end = parse_year(raw_to)
    if start and end and start > end:
        start, end = end, start
    return start, end


def reconstruct_inverted_abstract(index: Any) -> str:
    if not isinstance(index, dict):
        return ""
    positioned: list[tuple[int, str]] = []
    try:
        for token, positions in index.items():
            if isinstance(positions, list):
                for position in positions:
                    positioned.append((int(position), str(token)))
    except (TypeError, ValueError):
        return ""
    positioned.sort(key=lambda item: item[0])
    return clean_text(" ".join(token for _, token in positioned), limit=12000)


def make_result(
    *,
    source: str,
    title: Any,
    authors: Any = "",
    year: Any = None,
    abstract: Any = "",
    url: Any = "",
    doi: Any = "",
    venue: Any = "",
    citations: Any = None,
    publication_type: Any = "",
    open_access: Any = None,
    issn: Any = "",
    mean_citedness: Any = None,
    publication_date: Any = "",
) -> dict[str, Any]:
    normalized_doi = normalize_doi(doi)
    return {
        "source": source,
        "sources": [source],
        "title": clean_text(title, limit=800),
        "authors": clean_text(authors, limit=1600),
        "year": parse_int(year),
        "publication_date": clean_text(publication_date, limit=40),
        "abstract": clean_text(abstract, limit=12000),
        "url": safe_url(url) or doi_url(normalized_doi),
        "doi": normalized_doi,
        "venue": clean_text(venue, limit=500),
        "citations": parse_int(citations, 0) or 0,
        "publication_type": clean_text(publication_type, limit=120),
        "open_access": bool(open_access) if open_access is not None else None,
        "issn": clean_text(issn, limit=120),
        "journal_2yr_mean_citedness": parse_float(mean_citedness),
        "impact_factor": None,
        "impact_factor_source": "未配置",
        "summary_zh": "",
        "summary_status": "not_requested",
    }


def search_openalex(request: dict[str, Any]) -> list[dict[str, Any]]:
    keywords = request["keywords"]
    params: dict[str, str] = {
        "search": build_primary_query(request["query_terms"], request["match_mode"]),
        "per-page": str(request["fetch_limit"]),
        "sort": "relevance_score:desc",
    }
    filters: list[str] = []
    if request.get("from_year"):
        filters.append(f"from_publication_date:{request['from_year']}-01-01")
    if request.get("to_year"):
        filters.append(f"to_publication_date:{request['to_year']}-12-31")
    if filters:
        params["filter"] = ",".join(filters)
    email = os.environ.get("OPENALEX_EMAIL") or os.environ.get("CROSSREF_MAILTO")
    if email:
        params["mailto"] = email
    api_key = os.environ.get("OPENALEX_API_KEY", "").strip()
    headers = {"Authorization": f"Bearer {api_key}"} if api_key else None
    data = request_json("https://api.openalex.org/works?" + urllib.parse.urlencode(params), headers=headers)
    results: list[dict[str, Any]] = []
    for item in data.get("results", []) if isinstance(data, dict) else []:
        location = item.get("primary_location") or {}
        source_info = location.get("source") or {}
        authors = [
            as_text((authorship.get("author") or {}).get("display_name"))
            for authorship in (item.get("authorships") or [])
        ]
        abstract = reconstruct_inverted_abstract(item.get("abstract_inverted_index"))
        doi = item.get("doi") or (item.get("ids") or {}).get("doi")
        landing = source_info.get("homepage_url") or location.get("landing_page_url")
        result = make_result(
            source="OpenAlex",
            title=item.get("display_name"),
            authors="; ".join(author for author in authors if author),
            year=item.get("publication_year"),
            publication_date=item.get("publication_date"),
            abstract=abstract,
            url=item.get("doi") or landing or item.get("id"),
            doi=doi,
            venue=source_info.get("display_name"),
            citations=item.get("cited_by_count"),
            publication_type=item.get("type_crossref") or item.get("type"),
            open_access=(item.get("open_access") or {}).get("is_oa"),
            issn=source_info.get("issn_l") or next(iter(source_info.get("issns") or []), ""),
            mean_citedness=(item.get("summary_stats") or {}).get("2yr_mean_citedness"),
        )
        if result["title"]:
            results.append(result)
    return results


def crossref_query(keywords: list[str], match_mode: str) -> str:
    if match_mode == "all":
        return " ".join(f'"{keyword}"' for keyword in keywords)
    return " OR ".join(f'"{keyword}"' for keyword in keywords)


def search_crossref(request: dict[str, Any]) -> list[dict[str, Any]]:
    params: dict[str, str] = {
        "query.bibliographic": build_primary_query(request["query_terms"], request["match_mode"]),
        "rows": str(request["fetch_limit"]),
        "sort": "relevance",
        "order": "desc",
        "select": (
            "DOI,title,author,abstract,URL,published,issued,container-title,"
            "is-referenced-by-count,type,ISSN,subject,license,link"
        ),
    }
    filters: list[str] = []
    if request.get("from_year"):
        filters.append(f"from-pub-date:{request['from_year']}-01-01")
    if request.get("to_year"):
        filters.append(f"until-pub-date:{request['to_year']}-12-31")
    if filters:
        params["filter"] = ",".join(filters)
    email = os.environ.get("CROSSREF_MAILTO") or os.environ.get("OPENALEX_EMAIL")
    if email:
        params["mailto"] = email
    data = request_json("https://api.crossref.org/works?" + urllib.parse.urlencode(params))
    items = ((data or {}).get("message") or {}).get("items") or []
    results: list[dict[str, Any]] = []
    for item in items:
        authors = item.get("author") or []
        author_names: list[str] = []
        for author in authors:
            full_name = " ".join(
                part for part in [as_text(author.get("given")), as_text(author.get("family"))] if part
            ).strip()
            if not full_name:
                full_name = as_text(author.get("name"))
            if full_name:
                author_names.append(full_name)
        date_parts = ((item.get("published") or item.get("issued") or {}).get("date-parts") or [[]])[0]
        publication_year = parse_int(date_parts[0]) if date_parts else None
        date_values = [parse_int(part) for part in date_parts[:3]] if date_parts else []
        publication_date = (
            "-".join(str(value) if index == 0 else f"{value:02d}" for index, value in enumerate(date_values) if value is not None)
            if publication_year is not None else ""
        )
        doi = item.get("DOI")
        result = make_result(
            source="Crossref",
            title=first(item.get("title")),
            authors="; ".join(author_names),
            year=publication_year,
            publication_date=publication_date,
            abstract=item.get("abstract"),
            url=item.get("URL") or doi_url(normalize_doi(doi)),
            doi=doi,
            venue=first(item.get("container-title")),
            citations=item.get("is-referenced-by-count"),
            publication_type=item.get("type"),
            open_access=bool(item.get("license")),
            issn=first(item.get("ISSN")),
        )
        if result["title"]:
            results.append(result)
    return results


def search_arxiv(request: dict[str, Any]) -> list[dict[str, Any]]:
    group_parts: list[str] = []
    for group in request["query_terms"]:
        parts = [
            f'all:"{variant.replace(chr(34), chr(92) + chr(34))}"'
            for variant in group_variants(group)
        ]
        if parts:
            group_parts.append("(" + " OR ".join(parts) + ")")
    joiner = " AND " if request["match_mode"] == "all" else " OR "
    search_query = joiner.join(group_parts)
    start_year = request.get("from_year") or 1991
    end_year = request.get("to_year") or date.today().year
    search_query += f" AND submittedDate:[{start_year}01010000 TO {end_year}12312359]"
    params = {
        "search_query": search_query,
        "start": "0",
        "max_results": str(request["fetch_limit"]),
        "sortBy": "relevance",
        "sortOrder": "descending",
    }
    raw = request_bytes("https://export.arxiv.org/api/query?" + urllib.parse.urlencode(params))
    root = ET.fromstring(raw)
    ns = {
        "a": "http://www.w3.org/2005/Atom",
        "arxiv": "http://arxiv.org/schemas/atom",
    }
    results: list[dict[str, Any]] = []
    for entry in root.findall("a:entry", ns):
        entry_id = as_text(entry.findtext("a:id", default="", namespaces=ns))
        title = entry.findtext("a:title", default="", namespaces=ns)
        summary = entry.findtext("a:summary", default="", namespaces=ns)
        published = entry.findtext("a:published", default="", namespaces=ns)
        authors = [as_text(node.findtext("a:name", default="", namespaces=ns)) for node in entry.findall("a:author", ns)]
        journal_ref = entry.findtext("arxiv:journal_ref", default="", namespaces=ns)
        doi = entry.findtext("arxiv:doi", default="", namespaces=ns)
        link = ""
        for node in entry.findall("a:link", ns):
            if node.attrib.get("rel") == "alternate" or node.attrib.get("title") == "pdf":
                link = node.attrib.get("href", "")
                if node.attrib.get("rel") == "alternate":
                    break
        result = make_result(
            source="arXiv",
            title=title,
            authors="; ".join(author for author in authors if author),
            year=parse_year(published),
            publication_date=published[:10],
            abstract=summary,
            url=link or entry_id,
            doi=doi,
            venue=journal_ref or "arXiv",
            citations=0,
            publication_type="preprint",
            open_access=True,
        )
        if result["title"]:
            results.append(result)
    return results


def search_semantic_scholar(request: dict[str, Any]) -> list[dict[str, Any]]:
    params: dict[str, str] = {
        "query": build_flat_query(request["query_terms"], request["match_mode"]),
        "limit": str(request["fetch_limit"]),
        "fields": (
            "title,abstract,url,year,venue,authors,citationCount,"
            "influentialCitationCount,externalIds,publicationTypes,openAccessPdf,publicationDate"
        ),
    }
    if request.get("from_year") or request.get("to_year"):
        start = request.get("from_year") or ""
        end = request.get("to_year") or ""
        if start and end:
            params["year"] = f"{start}-{end}"
        elif start:
            params["year"] = f"{start}-"
        else:
            params["year"] = f"-{end}"
    headers: dict[str, str] = {}
    api_key = os.environ.get("SEMANTIC_SCHOLAR_API_KEY")
    if api_key:
        headers["x-api-key"] = api_key
    data = request_json(
        "https://api.semanticscholar.org/graph/v1/paper/search?" + urllib.parse.urlencode(params),
        headers=headers,
    )
    results: list[dict[str, Any]] = []
    for item in (data or {}).get("data", []) if isinstance(data, dict) else []:
        authors = [as_text(author.get("name")) for author in (item.get("authors") or [])]
        external = item.get("externalIds") or {}
        open_pdf = (item.get("openAccessPdf") or {}).get("url")
        publication_types = item.get("publicationTypes") or []
        result = make_result(
            source="Semantic Scholar",
            title=item.get("title"),
            authors="; ".join(author for author in authors if author),
            year=item.get("year"),
            publication_date=item.get("publicationDate"),
            abstract=item.get("abstract"),
            url=item.get("url") or open_pdf,
            doi=external.get("DOI"),
            venue=item.get("venue"),
            citations=item.get("citationCount"),
            publication_type=", ".join(as_text(value) for value in publication_types),
            open_access=bool(open_pdf),
        )
        if result["title"]:
            results.append(result)
    return results


def pubmed_query(request: dict[str, Any]) -> str:
    group_parts: list[str] = []
    for group in request["query_terms"]:
        parts = [f'"{variant}"[Title/Abstract]' for variant in group_variants(group)]
        if parts:
            group_parts.append("(" + " OR ".join(parts) + ")")
    joiner = " AND " if request["match_mode"] == "all" else " OR "
    query = joiner.join(group_parts)
    from_year = request.get("from_year")
    to_year = request.get("to_year")
    if from_year or to_year:
        start = f"{from_year}/01/01" if from_year else "1800/01/01"
        end = f"{to_year}/12/31" if to_year else f"{date.today().year}/12/31"
        query += f' AND ("{start}"[Date - Publication] : "{end}"[Date - Publication])'
    return query


def text_with_children(node: ET.Element | None) -> str:
    if node is None:
        return ""
    return as_text("".join(node.itertext()))


def search_pubmed(request: dict[str, Any]) -> list[dict[str, Any]]:
    common: dict[str, str] = {
        "db": "pubmed",
        "term": pubmed_query(request),
        "tool": "paper-search-web",
    }
    ncbi_email = os.environ.get("NCBI_EMAIL")
    ncbi_key = os.environ.get("NCBI_API_KEY")
    if ncbi_email:
        common["email"] = ncbi_email
    if ncbi_key:
        common["api_key"] = ncbi_key
    search_params = dict(common)
    search_params.update(
        {
            "retmode": "json",
            "retmax": str(request["fetch_limit"]),
            "sort": "relevance",
        }
    )
    search_data = request_json(
        "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi?"
        + urllib.parse.urlencode(search_params)
    )
    ids = ((search_data or {}).get("esearchresult") or {}).get("idlist") or []
    if not ids:
        return []
    fetch_params = dict(common)
    fetch_params.update({"retmode": "xml", "id": ",".join(ids)})
    raw = request_bytes(
        "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi?"
        + urllib.parse.urlencode(fetch_params)
    )
    root = ET.fromstring(raw)
    results: list[dict[str, Any]] = []
    for article in root.findall(".//PubmedArticle"):
        citation = article.find("MedlineCitation")
        article_node = citation.find("Article") if citation is not None else None
        if citation is None or article_node is None:
            continue
        pmid = text_with_children(citation.find("PMID"))
        title = text_with_children(article_node.find("ArticleTitle"))
        journal = text_with_children(article_node.find("Journal/Title"))
        issn = text_with_children(article_node.find("Journal/ISSN"))
        authors: list[str] = []
        for author in article_node.findall(".//AuthorList/Author"):
            collective = text_with_children(author.find("CollectiveName"))
            if collective:
                authors.append(collective)
                continue
            full_name = " ".join(
                part
                for part in [
                    text_with_children(author.find("ForeName")),
                    text_with_children(author.find("LastName")),
                ]
                if part
            ).strip()
            if full_name:
                authors.append(full_name)
        abstract_parts: list[str] = []
        for node in article_node.findall(".//Abstract/AbstractText"):
            label = node.attrib.get("Label") or node.attrib.get("NlmCategory") or ""
            value = text_with_children(node)
            if value:
                abstract_parts.append(f"{label}: {value}" if label else value)
        abstract = " ".join(abstract_parts)
        pub_date = article_node.find("Journal/JournalIssue/PubDate")
        year = text_with_children(pub_date.find("Year")) if pub_date is not None else ""
        if not year and pub_date is not None:
            year = parse_year(text_with_children(pub_date.find("MedlineDate")))
        publication_date = "-".join(
            part
            for part in [
                str(year or ""),
                text_with_children(pub_date.find("Month")) if pub_date is not None else "",
                text_with_children(pub_date.find("Day")) if pub_date is not None else "",
            ]
            if part
        )
        doi = ""
        pmcid = ""
        for identifier in citation.findall(".//ArticleId"):
            id_type = identifier.attrib.get("IdType", "").lower()
            if id_type == "doi":
                doi = text_with_children(identifier)
            elif id_type == "pmc":
                pmcid = text_with_children(identifier)
        publication_types = [
            text_with_children(node)
            for node in article_node.findall(".//PublicationTypeList/PublicationType")
        ]
        result = make_result(
            source="PubMed",
            title=title,
            authors="; ".join(authors),
            year=parse_year(year),
            publication_date=publication_date,
            abstract=abstract,
            url=f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/" if pmid else doi_url(normalize_doi(doi)),
            doi=doi,
            venue=journal,
            citations=0,
            publication_type=", ".join(publication_types),
            open_access=bool(pmcid),
            issn=issn,
        )
        if result["title"]:
            results.append(result)
    return results


def search_europe_pmc(request: dict[str, Any]) -> list[dict[str, Any]]:
    group_parts: list[str] = []
    for group in request["query_terms"]:
        parts = [
            f'TITLE_ABS:"{variant.replace(chr(34), chr(92) + chr(34))}"'
            for variant in group_variants(group)
        ]
        if parts:
            group_parts.append("(" + " OR ".join(parts) + ")")
    joiner = " AND " if request["match_mode"] == "all" else " OR "
    query = joiner.join(group_parts)
    if request.get("from_year") or request.get("to_year"):
        start = f"{request.get('from_year') or 1800}-01-01"
        end = f"{request.get('to_year') or date.today().year}-12-31"
        query += f" AND (FIRST_PDATE:[{start} TO {end}])"
    params = {
        "query": query,
        "format": "json",
        "pageSize": str(request["fetch_limit"]),
        "resultType": "core",
    }
    data = request_json(
        "https://www.ebi.ac.uk/europepmc/webservices/rest/search?"
        + urllib.parse.urlencode(params)
    )
    items = ((data or {}).get("resultList") or {}).get("result") or []
    results: list[dict[str, Any]] = []
    for item in items:
        doi = item.get("doi")
        source_id = as_text(item.get("source"))
        identifier = as_text(item.get("id"))
        url = doi_url(normalize_doi(doi))
        if not url and identifier:
            url = f"https://europepmc.org/article/{source_id or 'MED'}/{identifier}"
        journal_info = item.get("journalInfo") or {}
        journal = journal_info.get("journal") or {}
        pub_types = ((item.get("pubTypeList") or {}).get("pubType") or [])
        result = make_result(
            source="Europe PMC",
            title=item.get("title"),
            authors=item.get("authorString"),
            year=item.get("pubYear"),
            publication_date=item.get("firstPublicationDate"),
            abstract=item.get("abstractText"),
            url=url,
            doi=doi,
            venue=item.get("journalTitle") or journal.get("title"),
            citations=item.get("citedByCount"),
            publication_type=", ".join(as_text(value) for value in pub_types),
            open_access=str(item.get("isOpenAccess", "")).lower() == "y" or bool(item.get("pmcid")),
            issn=item.get("journalIssn") or journal.get("issn"),
        )
        if result["title"]:
            results.append(result)
    return results


SOURCE_DEFINITIONS: list[dict[str, Any]] = [
    {
        "id": "openalex",
        "name": "OpenAlex",
        "adapter": search_openalex,
        "description": "综合学术图谱，覆盖期刊、会议、书籍与预印本。",
        "homepage": "https://openalex.org",
        "capabilities": ["摘要", "被引数", "开放获取"],
    },
    {
        "id": "crossref",
        "name": "Crossref",
        "adapter": search_crossref,
        "description": "出版机构元数据网络，DOI、期刊和作者信息可靠。",
        "homepage": "https://www.crossref.org",
        "capabilities": ["DOI", "被引数", "期刊信息"],
    },
    {
        "id": "arxiv",
        "name": "arXiv",
        "adapter": search_arxiv,
        "description": "物理、数学、计算机等领域的开放预印本。",
        "homepage": "https://arxiv.org",
        "capabilities": ["摘要", "免费全文", "预印本"],
    },
    {
        "id": "semantic_scholar",
        "name": "Semantic Scholar",
        "adapter": search_semantic_scholar,
        "description": "AI 驱动的学术搜索，适合查看引用和开放获取链接。",
        "homepage": "https://www.semanticscholar.org",
        "capabilities": ["摘要", "被引数", "开放获取"],
    },
    {
        "id": "pubmed",
        "name": "PubMed",
        "adapter": search_pubmed,
        "description": "生物医学和生命科学领域的权威文献数据库。",
        "homepage": "https://pubmed.ncbi.nlm.nih.gov",
        "capabilities": ["摘要", "医学主题词", "开放获取"],
    },
    {
        "id": "europe_pmc",
        "name": "Europe PMC",
        "adapter": search_europe_pmc,
        "description": "生命科学开放文献库，覆盖 PubMed、PMC 和预印本。",
        "homepage": "https://europepmc.org",
        "capabilities": ["摘要", "被引数", "全文链接"],
    },
]
SOURCE_BY_ID = {source["id"]: source for source in SOURCE_DEFINITIONS}


def source_public_info(source: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": source["id"],
        "name": source["name"],
        "description": source["description"],
        "homepage": source["homepage"],
        "capabilities": source["capabilities"],
    }


def merge_records(base: dict[str, Any], incoming: dict[str, Any]) -> dict[str, Any]:
    merged = dict(base)
    for field in ("title", "abstract", "authors", "venue", "publication_type", "publication_date", "issn"):
        current = merged.get(field) or ""
        candidate = incoming.get(field) or ""
        if len(candidate) > len(current):
            merged[field] = candidate
    for field in ("year", "citations", "impact_factor", "journal_2yr_mean_citedness", "open_access"):
        current = merged.get(field)
        candidate = incoming.get(field)
        if field in {"citations", "journal_2yr_mean_citedness"}:
            values = [value for value in (current, candidate) if value is not None]
            merged[field] = max(values) if values else None
        elif current is None and candidate is not None:
            merged[field] = candidate
    merged["sources"] = list(dict.fromkeys((base.get("sources") or []) + (incoming.get("sources") or [])))
    if not merged.get("url"):
        merged["url"] = incoming.get("url") or ""
    if not merged.get("doi"):
        merged["doi"] = incoming.get("doi") or ""
    return merged


def deduplicate_results(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    unique: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for item in items:
        if not item.get("title"):
            continue
        doi = normalize_doi(item.get("doi"))
        title_key = normalize_title(item.get("title"))
        key = f"doi:{doi}" if doi else f"title:{title_key}"
        if not title_key and not doi:
            continue
        if key not in unique:
            unique[key] = item
            order.append(key)
        else:
            unique[key] = merge_records(unique[key], item)
    results: list[dict[str, Any]] = []
    for key in order:
        item = unique[key]
        item["source"] = " / ".join(item.get("sources") or [])
        results.append(item)
    results.sort(
        key=lambda item: (
            -(item.get("year") or 0),
            -(item.get("citations") or 0),
            normalize_title(item.get("title")),
        )
    )
    return results


def enrich_metrics(results: list[dict[str, Any]]) -> None:
    for item in results:
        impact, impact_source = JOURNAL_METRICS.lookup(item.get("venue"), item.get("issn"))
        item["impact_factor"] = impact
        item["impact_factor_source"] = impact_source


def normalize_search_request(payload: dict[str, Any]) -> dict[str, Any]:
    keywords = split_keywords(payload.get("keywords"))
    if not keywords:
        raise ValueError("请至少输入一个关键词。")
    query_terms = TRANSLATOR.expand(keywords)
    deepseek_term_selection = None
    deepseek_term_warnings: list[str] = []
    use_deepseek_terms = payload.get("use_deepseek_terms", True) is not False
    if DEEPSEEK.configured and use_deepseek_terms:
        try:
            deepseek_term_selection = DEEPSEEK.select_search_terms(keywords)
            query_terms = merge_deepseek_terms(query_terms, deepseek_term_selection)
        except Exception as exc:
            deepseek_term_warnings.append(str(exc))
    selected = payload.get("sources")
    if not isinstance(selected, list):
        selected = []
    selected = [str(source_id) for source_id in selected if str(source_id) in SOURCE_BY_ID]
    if not selected:
        selected = ["openalex", "crossref"]
    match_mode = str(payload.get("match_mode") or "all").lower()
    if match_mode not in {"all", "any"}:
        match_mode = "all"
    from_year, to_year = normalize_year_range(payload.get("from_year"), payload.get("to_year"))
    max_results = clamp_max_results(payload.get("max_results"))
    match_scope = str(payload.get("match_scope") or "title_abstract")
    if match_scope not in {"title_abstract", "title"}:
        raise ValueError("匹配范围必须为标题与摘要，或仅标题。")
    fetch_limit = min(MAX_RESULTS, max(max_results * 10, 50))
    warnings = [
        warning
        for group in query_terms
        for warning in group.get("translation_warnings") or []
        if warning
    ]
    return {
        "keywords": keywords,
        "query_terms": query_terms,
        "has_chinese_keywords": any(any("\u3400" <= char <= "\u9fff" for char in keyword) for keyword in keywords),
        "translation_warnings": warnings,
        "translation_provider": TRANSLATOR.provider,
        "deepseek_terms_requested": use_deepseek_terms,
        "deepseek_terms_used": bool(deepseek_term_selection),
        "deepseek_term_selection": deepseek_term_selection,
        "deepseek_term_warnings": deepseek_term_warnings,
        "sources": selected,
        "match_mode": match_mode,
        "match_scope": match_scope,
        "exclude_keywords": split_keywords(payload.get("exclude_keywords")),
        "max_results": max_results,
        "fetch_limit": fetch_limit,
        "from_year": from_year,
        "to_year": to_year,
    }


def normalize_search_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", as_text(value)).casefold()
    text = re.sub(r"[^\w\u3400-\u9fff]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def variant_in_text(haystack: str, variant: str) -> bool:
    if not haystack or not variant:
        return False
    if re.fullmatch(r"[a-z0-9][a-z0-9 ._-]*", variant):
        pattern = r"(?<![a-z0-9])" + re.escape(variant) + r"(?![a-z0-9])"
        return re.search(pattern, haystack) is not None
    return variant in haystack


def merge_deepseek_terms(
    query_terms: list[dict[str, Any]],
    selection: dict[str, Any] | None,
) -> list[dict[str, Any]]:
    if not selection:
        return query_terms
    by_keyword = {
        normalize_search_text(group.get("keyword")): group
        for group in query_terms
        if group.get("keyword")
    }
    for selected in selection.get("groups") or []:
        if not isinstance(selected, dict):
            continue
        target = by_keyword.get(normalize_search_text(selected.get("keyword")))
        if not target:
            continue
        variants = target.setdefault("variants", [])
        variant_sources = target.setdefault("variant_sources", {})
        for variant in group_variants(target):
            if variant == target.get("keyword"):
                variant_sources.setdefault(variant, ["输入词"])
            else:
                variant_sources.setdefault(
                    variant,
                    list(target.get("translation_sources") or ["翻译扩展"]),
                )
        for term in selected.get("terms") or []:
            term = clean_text(term, limit=90)
            if not term:
                continue
            if TRANSLATOR.is_component_only(str(target.get("keyword") or ""), term):
                target.setdefault("component_filtered_terms", []).append(term)
                continue
            normalized = normalize_search_text(term)
            existing = next(
                (variant for variant in variants if normalize_search_text(variant) == normalized),
                term,
            )
            add_variants(variants, [existing])
            variant_sources[existing] = list(
                dict.fromkeys((variant_sources.get(existing) or []) + ["DeepSeek"])
            )
        target["translation_sources"] = list(
            dict.fromkeys((target.get("translation_sources") or []) + ["DeepSeek"])
        )
    return query_terms


def relevance_filter(
    items: list[dict[str, Any]],
    query_terms: list[dict[str, Any]],
    match_mode: str,
    match_scope: str = "title_abstract",
    exclude_keywords: list[str] | None = None,
) -> list[dict[str, Any]]:
    """按可见标题和摘要校验概念，保留依据；零命中绝不回退到候选集。"""
    if not query_terms:
        return items
    strict: list[dict[str, Any]] = []
    exclusions = [normalize_search_text(value) for value in (exclude_keywords or [])]
    for item in items:
        title = normalize_search_text(item.get("title"))
        abstract = normalize_search_text(item.get("abstract"))
        if any(variant_in_text(title, word) or variant_in_text(abstract, word) for word in exclusions if word):
            continue
        group_matches: list[bool] = []
        evidence: list[dict[str, Any]] = []
        score = 0
        for group in query_terms:
            variants = [(value, normalize_search_text(value)) for value in group_variants(group)]
            title_terms = [raw for raw, value in variants if value and variant_in_text(title, value)]
            abstract_terms = [raw for raw, value in variants if value and variant_in_text(abstract, value)] if match_scope != "title" else []
            group_matches.append(bool(title_terms or abstract_terms))
            if title_terms or abstract_terms:
                score += 3 if title_terms else 1
                evidence.append({
                    "keyword": group.get("keyword", ""),
                    "field": "标题" if title_terms else "摘要",
                    "terms": title_terms or abstract_terms,
                })
        matched = all(group_matches) if match_mode == "all" else any(group_matches)
        if matched:
            strict.append({**item, "match_evidence": evidence, "relevance_score": round(100 * score / (3 * len(query_terms)), 1)})
    strict.sort(key=lambda item: (-item["relevance_score"], -(item.get("year") or 0)))
    return strict


def search_all_sources(payload: dict[str, Any]) -> dict[str, Any]:
    request = normalize_search_request(payload)
    statuses: list[dict[str, Any] | None] = [None] * len(request["sources"])
    raw_results: list[dict[str, Any]] = []
    started = time.perf_counter()
    with ThreadPoolExecutor(max_workers=min(8, len(request["sources"]))) as executor:
        future_map = {}
        for index, source_id in enumerate(request["sources"]):
            source = SOURCE_BY_ID[source_id]
            source_started = time.perf_counter()
            future = executor.submit(source["adapter"], request)
            future_map[future] = (index, source, source_started)
        for future in as_completed(future_map):
            index, source, source_started = future_map[future]
            try:
                items = future.result()
                if not isinstance(items, list):
                    items = []
                candidate_count = len(items)
                items = relevance_filter(items, request["query_terms"], request["match_mode"], request["match_scope"], request["exclude_keywords"])
                items = items[: request["max_results"]]
                raw_results.extend(items)
                message = f"获取到 {len(items)} 条记录"
                if len(items) != candidate_count:
                    message += f"（候选 {candidate_count} 条，已按相关性筛选）"
                statuses[index] = {
                    "id": source["id"],
                    "name": source["name"],
                    "status": "ok",
                    "count": len(items),
                    "elapsed_ms": int((time.perf_counter() - source_started) * 1000),
                    "message": message,
                }
            except Exception as exc:
                message = clean_text(str(exc), limit=300) or exc.__class__.__name__
                statuses[index] = {
                    "id": source["id"],
                    "name": source["name"],
                    "status": "error",
                    "count": 0,
                    "elapsed_ms": int((time.perf_counter() - source_started) * 1000),
                    "message": message,
                }
    results = deduplicate_results(raw_results)
    # 合并来源可能补齐摘要，重新计算最终可见记录的命中依据与排序。
    results = relevance_filter(results, request["query_terms"], request["match_mode"], request["match_scope"], request["exclude_keywords"])
    enrich_metrics(results)
    clean_statuses = [status for status in statuses if status is not None]
    return {
        "query": request,
        "results": results,
        "sources": clean_statuses,
        "total": len(results),
        "raw_total": len(raw_results),
        "elapsed_ms": int((time.perf_counter() - started) * 1000),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }


def format_query_expansion(query: dict[str, Any]) -> str:
    parts: list[str] = []
    for group in query.get("query_terms") or []:
        keyword = clean_text(group.get("keyword"))
        variant_sources = group.get("variant_sources") or {}
        variants: list[str] = []
        for raw_value in group_variants(group):
            value = clean_text(raw_value)
            if not value:
                continue
            sources = variant_sources.get(raw_value) or []
            label = f"{value} [+{'/'.join(sources)}]" if sources else value
            variants.append(label)
        if keyword and variants:
            parts.append(f"{keyword} → {' | '.join(variants)}")
    return "；".join(parts)


def format_deepseek_terms(query: dict[str, Any]) -> str:
    terms: list[str] = []
    for group in query.get("query_terms") or []:
        variant_sources = group.get("variant_sources") or {}
        for value in group_variants(group):
            sources = variant_sources.get(value) or []
            if "DeepSeek" in sources:
                add_variants(terms, [value])
    return "；".join(terms)


def excel_safe(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def create_excel_bytes(
    results: list[dict[str, Any]],
    query: dict[str, Any],
    source_statuses: list[dict[str, Any]] | None = None,
) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:
        raise RuntimeError("导出 Excel 需要 openpyxl，请运行: pip install openpyxl") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "论文汇总"
    headers = [
        ("序号", 7),
        ("论文名称", 48),
        ("摘要", 75),
        ("AI摘要概括", 60),
        ("作者", 32),
        ("年份", 8),
        ("期刊/会议", 28),
        ("来源", 22),
        ("期刊2年平均被引", 18),
        ("论文被引数", 12),
        ("DOI", 28),
        ("类型", 18),
        ("开放获取", 12),
        ("检索关键词", 24),
        ("检索词扩展", 58),
        ("DeepSeek选取词", 50),
        ("检索时间", 20),
        ("网址", 34),
    ]
    header_fill = PatternFill("solid", fgColor="17324D")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, (title, width) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R{max(1, len(results) + 1)}"

    keyword_text = "；".join(query.get("keywords") or [])
    expansion_text = format_query_expansion(query)
    deepseek_terms_text = format_deepseek_terms(query)
    match_label = "全部关键词" if query.get("match_mode") == "all" else "任一关键词"
    retrieved_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    highlight_variants = [
        value
        for value in all_variant_terms(query.get("query_terms") or [])
        if len(value) >= 2
    ]
    highlight_variants.sort(key=len, reverse=True)

    def highlighted_title(value: Any) -> Any:
        text_value = clean_text(value)
        if not text_value or not highlight_variants:
            return text_value
        def title_pattern(variant: str) -> str:
            escaped = re.escape(variant)
            if re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9 ._-]*", variant):
                return r"(?<![A-Za-z0-9])" + escaped + r"(?![A-Za-z0-9])"
            return escaped

        pattern = re.compile(
            "(" + "|".join(title_pattern(variant) for variant in highlight_variants) + ")",
            re.IGNORECASE,
        )
        matches = list(pattern.finditer(text_value))
        if not matches:
            return text_value
        rich_text = CellRichText()
        position = 0
        for match in matches:
            if match.start() > position:
                rich_text.append(text_value[position:match.start()])
            rich_text.append(TextBlock(InlineFont(color="FF0000", b=True), match.group(0)))
            position = match.end()
        if position < len(text_value):
            rich_text.append(text_value[position:])
        return rich_text

    for row_index, item in enumerate(results, start=2):
        values = [
            row_index - 1,
            highlighted_title(item.get("title") or ""),
            item.get("abstract") or "",
            item.get("summary_zh") or "",
            item.get("authors") or "",
            item.get("year") or "",
            item.get("venue") or "",
            item.get("source") or "",
            item.get("journal_2yr_mean_citedness"),
            item.get("citations") or 0,
            item.get("doi") or "",
            item.get("publication_type") or "",
            "是" if item.get("open_access") is True else ("否" if item.get("open_access") is False else "未知"),
            keyword_text,
            expansion_text,
            deepseek_terms_text,
            retrieved_at,
            item.get("url") or "",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column)
            cell.value = value if column == 2 and isinstance(value, CellRichText) else excel_safe(value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=column in {2, 3, 4, 5, 7, 8, 11, 12, 14, 15, 16})
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F6F9FC")
        sheet.row_dimensions[row_index].height = 54 if item.get("abstract") else 30
        mean_cell = sheet.cell(row=row_index, column=9)
        mean_cell.number_format = "0.000"
        url_cell = sheet.cell(row=row_index, column=18)
        if isinstance(url_cell.value, str) and url_cell.value.startswith(("http://", "https://")):
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    status_sheet = workbook.create_sheet("来源状态")
    status_headers = ["数据源", "状态", "原始条数", "耗时(ms)", "说明"]
    for column, title in enumerate(status_headers, start=1):
        cell = status_sheet.cell(row=1, column=column, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, status in enumerate(source_statuses or [], start=2):
        row_values = [
            status.get("name") or status.get("id") or "",
            "成功" if status.get("status") == "ok" else "失败",
            status.get("count") or 0,
            status.get("elapsed_ms") or 0,
            status.get("message") or "",
        ]
        for column, value in enumerate(row_values, start=1):
            status_sheet.cell(row=row_index, column=column, value=excel_safe(value)).border = border
    for column, width in zip("ABCDE", [22, 10, 12, 12, 70]):
        status_sheet.column_dimensions[column].width = width
    status_sheet.freeze_panes = "A2"
    status_sheet.auto_filter.ref = f"A1:E{max(1, len(source_statuses or []) + 1)}"

    note_sheet = workbook.create_sheet("检索说明")
    note_rows = [
        ("项目", "内容"),
        ("检索关键词", keyword_text),
        ("检索词扩展", expansion_text),
        ("翻译方式", "中文原词 + 本地学术词库 + MyMemory 在线翻译（可在环境变量中调整）"),
        ("匹配方式", match_label),
        ("年份范围", f"{query.get('from_year') or '不限'} - {query.get('to_year') or '不限'}"),
        ("每个数据源上限", query.get("max_results") or ""),
        ("检索时间", retrieved_at),
        ("结果总数", len(results)),
        ("期刊2年平均被引说明", "来自 OpenAlex summary_stats.2yr_mean_citedness，仅作为期刊层面的辅助指标。"),
        ("DeepSeek选取词", deepseek_terms_text or "本次未使用 DeepSeek 选词"),
        ("AI摘要概括", "由 DeepSeek API 根据论文摘要生成，仅用于辅助阅读，请以论文原文为准。"),
        ("数据使用提醒", "结果来自公开 API，请核对原文和数据库使用条款；不要用本工具绕过付费墙或登录限制。"),
    ]
    for row_index, (label, value) in enumerate(note_rows, start=1):
        note_sheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        note_sheet.cell(row=row_index, column=2, value=excel_safe(value))
        note_sheet.cell(row=row_index, column=1).border = border
        note_sheet.cell(row=row_index, column=2).border = border
        note_sheet.cell(row=row_index, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    note_sheet.column_dimensions["A"].width = 22
    note_sheet.column_dimensions["B"].width = 90
    workbook.properties.title = "论文检索结果"
    workbook.properties.creator = "Paper Search Web"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class PaperSearchHandler(SimpleHTTPRequestHandler):
    server_version = f"PaperSearchWeb/{APP_VERSION}"

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, directory=str(STATIC_DIR), **kwargs)

    def log_message(self, fmt: str, *args: Any) -> None:
        print(f"[web] {self.address_string()} - {fmt % args}")

    def end_headers(self) -> None:
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        if self.path.startswith("/api/"):
            self.send_header("Cache-Control", "no-store")
        else:
            # Static assets revalidate on every load so UI updates reach browsers immediately.
            self.send_header("Cache-Control", "no-cache")
        super().end_headers()

    def send_json(self, payload: Any, status: int = 200) -> None:
        body = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_json(self) -> dict[str, Any]:
        length = parse_int(self.headers.get("Content-Length"), 0) or 0
        if length <= 0:
            return {}
        if length > 5_000_000:
            raise ValueError("请求体过大。")
        raw = self.rfile.read(length)
        try:
            payload = json.loads(raw.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("请求不是有效的 JSON。") from exc
        if not isinstance(payload, dict):
            raise ValueError("请求体必须是 JSON 对象。")
        return payload

    def do_GET(self) -> None:
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path == "/api/health":
            self.send_json(
                {
                    "ok": True,
                    "version": APP_VERSION,
                    "journal_metrics_count": JOURNAL_METRICS.count(),
                    "translation_dictionary_count": TRANSLATOR.dictionary_count(),
                    "translation_provider": TRANSLATOR.provider,
                    "deepseek_configured": DEEPSEEK.configured,
                    "deepseek_model": DEEPSEEK.model,
                    "sources": len(SOURCE_DEFINITIONS),
                }
            )
            return
        if parsed.path == "/api/sources":
            self.send_json(
                {
                    "sources": [source_public_info(source) for source in SOURCE_DEFINITIONS],
                    "journal_metrics_count": JOURNAL_METRICS.count(),
                    "translation_dictionary_count": TRANSLATOR.dictionary_count(),
                    "translation_provider": TRANSLATOR.provider,
                    "translation_note": (
                        "检测到中文关键词时，会自动保留中文并扩展英文学术译法和同义词，"
                        "然后使用中英文变体共同检索。"
                    ),
                    "deepseek": DEEPSEEK.status(),
                }
            )
            return
        if parsed.path == "/api/deepseek/status":
            self.send_json(DEEPSEEK.status())
            return
        if parsed.path == "/api/deepseek/config":
            self.send_json(DEEPSEEK.status())
            return
        if parsed.path == "/api/journal-metrics":
            self.send_json(
                {
                    "count": JOURNAL_METRICS.count(),
                    "file": "data/journal_metrics.csv",
                    "description": "journal,issn,eissn,impact_factor,year,source",
                }
            )
            return
        if parsed.path.startswith("/api/"):
            self.send_json({"error": "接口不存在"}, status=404)
            return
        super().do_GET()

    def do_POST(self) -> None:
        parsed = urllib.parse.urlparse(self.path)
        try:
            payload = self.read_json()
            if parsed.path == "/api/search":
                self.send_json(search_all_sources(payload))
                return
            if parsed.path == "/api/deepseek/config":
                action = str(payload.get("action") or "save").strip().lower()
                if action == "clear":
                    status = DEEPSEEK.clear()
                else:
                    api_key = str(payload.get("api_key") or "").strip()
                    model = str(payload.get("model") or "").strip() or None
                    base_url = str(payload.get("base_url") or "").strip() or None
                    status = DEEPSEEK.configure(api_key, model=model, base_url=base_url)
                self.send_json({"ok": True, **status})
                return
            if parsed.path == "/api/summarize":
                results = payload.get("results")
                if not isinstance(results, list):
                    raise ValueError("缺少需要概括的结果。")
                try:
                    self.send_json(DEEPSEEK.summarize_results(results))
                except DeepSeekError as exc:
                    self.send_json({"error": str(exc)}, status=400)
                return
            if parsed.path == "/api/export":
                results = payload.get("results")
                if not isinstance(results, list):
                    raise ValueError("缺少检索结果。")
                query = payload.get("query") if isinstance(payload.get("query"), dict) else {}
                statuses = payload.get("sources") if isinstance(payload.get("sources"), list) else []
                body = create_excel_bytes(results, query, statuses)
                keywords = query.get("keywords") if isinstance(query.get("keywords"), list) else []
                keyword_part = re.sub(r"[^\w\u3400-\u9fff-]+", "_", "_".join(str(value) for value in keywords))[:48]
                filename_ascii = f"paper_search_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
                filename_utf8 = f"论文检索_{keyword_part or 'results'}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
                self.send_response(200)
                self.send_header(
                    "Content-Type",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                self.send_header("Content-Length", str(len(body)))
                self.send_header(
                    "Content-Disposition",
                    f"attachment; filename=\"{filename_ascii}\"; filename*=UTF-8''{urllib.parse.quote(filename_utf8)}",
                )
                self.end_headers()
                self.wfile.write(body)
                return
            self.send_json({"error": "接口不存在"}, status=404)
        except ValueError as exc:
            self.send_json({"error": str(exc)}, status=400)
        except RuntimeError as exc:
            self.send_json({"error": str(exc)}, status=500)
        except Exception as exc:
            print(f"[error] {exc.__class__.__name__}: {exc}", file=sys.stderr)
            self.send_json({"error": f"服务器处理失败: {exc.__class__.__name__}"}, status=500)


def main() -> None:
    parser = argparse.ArgumentParser(description="Paper Search Web")
    parser.add_argument("--host", default=os.environ.get("PAPER_SEARCH_HOST", "127.0.0.1"))
    parser.add_argument("--port", type=int, default=int(os.environ.get("PAPER_SEARCH_PORT", "8000")))
    parser.add_argument("--open", action="store_true", help="启动后自动打开浏览器")
    args = parser.parse_args()
    server = ThreadingHTTPServer((args.host, args.port), PaperSearchHandler)
    url = f"http://{args.host}:{args.port}"
    print("=" * 58)
    print("Paper Search Web 已启动")
    print(f"访问地址: {url}")
    print("按 Ctrl+C 停止服务")
    print("=" * 58)
    if args.open:
        try:
            webbrowser.open(url)
        except Exception:
            pass
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n服务已停止。")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
